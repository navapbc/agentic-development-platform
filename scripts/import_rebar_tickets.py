"""Import a Jira xlsx export into the local rebar ticket store.

The importer maps a Jira issue export onto rebar tickets so a deployment can load
its own tracker data rather than any bundled sample. It requires a path to the
xlsx export. It maps issue type, status, and priority to the rebar vocabulary,
folds the Acceptance Criteria custom field into the description, preserves the
Jira issue key as import provenance, and carries comments across. It is
idempotent by Jira issue key, so a re-run skips tickets already imported.

Seeding the store is otherwise rebar's own capability. This importer covers the
xlsx-to-rebar path the platform drives directly.
"""

from __future__ import annotations

import argparse
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

READER_IDENTITY_NAME = "agent-platform"
READER_IDENTITY_EMAIL = "agent-platform@users.noreply.invalid"

# rebar restricts work types to bug, epic, story, and task. Unmapped Jira types
# coerce to task, with the original type preserved in the description.
_TYPE_MAP = {
    "task": "task",
    "bug": "bug",
    "story": "story",
    "epic": "epic",
    "spike": "task",
    "design task": "task",
    "sub-task": "task",
    "subtask": "task",
}

# rebar statuses are idea, open, in_progress, closed, and blocked. idea holds
# not-yet-actionable work and is excluded from rebar.ready.
_STATUS_MAP = {
    "to do": "open",
    "open": "open",
    "in progress": "in_progress",
    "in review": "in_progress",
    "done": "closed",
    "closed": "closed",
    "resolved": "closed",
    "blocked": "blocked",
    "backlog": "idea",
    "refinement": "idea",
}

# rebar priority is an integer 0-4 where 0 is the highest.
_PRIORITY_MAP = {
    "highest": 0,
    "high": 1,
    "medium": 2,
    "low": 3,
    "lowest": 4,
}
_DEFAULT_PRIORITY = 2


@dataclass
class ImportReport:
    created: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    transitioned: int = 0


def _git(root: Path, *args: str) -> None:
    subprocess.run(["git", *args], cwd=root, check=True, capture_output=True, text=True)


def ensure_store(repo_root: Path) -> None:
    import rebar

    repo_root.mkdir(parents=True, exist_ok=True)
    if not (repo_root / ".git").exists():
        _git(repo_root, "init", "-q", "-b", "main")
        _git(repo_root, "config", "user.name", READER_IDENTITY_NAME)
        _git(repo_root, "config", "user.email", READER_IDENTITY_EMAIL)
    rebar.init_repo(repo_root=str(repo_root), force_new_store=True)
    _ensure_reader_identity(repo_root)


def _ensure_reader_identity(repo_root: Path) -> None:
    import rebar

    if rebar.resolve_current_identity(repo_root=str(repo_root)):
        return
    identity_id = rebar.create_identity(
        READER_IDENTITY_NAME, READER_IDENTITY_EMAIL, repo_root=str(repo_root)
    )
    rebar.use_identity(str(identity_id), repo_root=str(repo_root))


def reset_store(repo_root: Path) -> None:
    """Wipe the local store and re-init an empty store with the reader identity."""
    import shutil

    if repo_root.exists():
        shutil.rmtree(repo_root)
    ensure_store(repo_root)


def _header_index(header: tuple[Any, ...]) -> dict[str, list[int]]:
    index: dict[str, list[int]] = {}
    for i, name in enumerate(header):
        if name is None:
            continue
        index.setdefault(str(name).strip(), []).append(i)
    return index


def _cell(row: tuple[Any, ...], index: dict[str, list[int]], name: str) -> str:
    positions = index.get(name)
    if not positions:
        return ""
    value = row[positions[0]]
    return "" if value is None else str(value).strip()


def _all_cells(row: tuple[Any, ...], index: dict[str, list[int]], name: str) -> list[str]:
    out: list[str] = []
    for pos in index.get(name, []):
        value = row[pos]
        if value is not None and str(value).strip():
            out.append(str(value).strip())
    return out


def _rebar_type(issue_type: str) -> str:
    return _TYPE_MAP.get(issue_type.strip().lower(), "task")


def _rebar_status(status: str) -> str:
    return _STATUS_MAP.get(status.strip().lower(), "open")


def _rebar_priority(priority: str) -> int:
    return _PRIORITY_MAP.get(priority.strip().lower(), _DEFAULT_PRIORITY)


def _compose_description(row: tuple[Any, ...], index: dict[str, list[int]]) -> str:
    parts: list[str] = []
    description = _cell(row, index, "Description")
    if description:
        parts.append(description)
    acceptance = _cell(row, index, "Custom field (Acceptance Criteria)")
    if acceptance:
        parts.append(f"Acceptance Criteria:\n{acceptance}")
    parent_key = _cell(row, index, "Parent key")
    parent_summary = _cell(row, index, "Parent summary")
    if parent_key:
        suffix = f" ({parent_summary})" if parent_summary else ""
        parts.append(f"Parent: {parent_key}{suffix}")
    points = _cell(row, index, "Custom field (Story Points)")
    if points:
        parts.append(f"Story Points: {points}")
    issue_type = _cell(row, index, "Issue Type")
    if issue_type and _rebar_type(issue_type) != issue_type.strip().lower():
        parts.append(f"Jira type: {issue_type}")
    return "\n\n".join(parts)


def _existing_source_keys(repo_root: Path) -> set[str]:
    import rebar

    keys: set[str] = set()
    for ticket in rebar.list_tickets(repo_root=str(repo_root), include_archived=True):
        source_id = ticket.get("source_id") or ""
        if source_id:
            keys.add(str(source_id))
    return keys


def import_xlsx(
    xlsx_path: Path, repo_root: Path, *, statuses: set[str] | None = None, limit: int = 0
) -> ImportReport:
    import openpyxl
    import rebar

    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows)
    index = _header_index(header)
    report = ImportReport()
    already = _existing_source_keys(repo_root)
    seen_titles = {
        str(t.get("title") or "")
        for t in rebar.list_tickets(repo_root=str(repo_root), include_archived=True)
    }

    for row in rows:
        if limit and len(report.created) >= limit:
            break
        summary = _cell(row, index, "Summary")
        issue_key = _cell(row, index, "Issue key")
        jira_status = _cell(row, index, "Status")
        if not summary:
            continue
        if statuses is not None and jira_status.strip().lower() not in statuses:
            continue
        if issue_key and issue_key in already:
            report.skipped.append(issue_key or summary)
            continue
        if not issue_key and summary in seen_titles:
            report.skipped.append(summary)
            continue

        source = {"source_id": issue_key, "source_env": "jira"} if issue_key else None
        ticket_id = rebar.create_ticket(
            _rebar_type(_cell(row, index, "Issue Type")),
            summary,
            description=_compose_description(row, index) or None,
            priority=_rebar_priority(_cell(row, index, "Priority")),
            tags=_all_cells(row, index, "Components") or None,
            source=source,
            repo_root=str(repo_root),
        )
        ticket_id = str(ticket_id)
        report.created.append(issue_key or ticket_id)

        for body in _all_cells(row, index, "Comment"):
            rebar.comment(ticket_id, body, repo_root=str(repo_root))

        target = _rebar_status(jira_status)
        if target != "open":
            try:
                rebar.transition(
                    ticket_id, "open", target, force=True, reason="jira import",
                    repo_root=str(repo_root),
                )
                report.transitioned += 1
            except rebar.RebarError:
                pass

    return report


def _default_repo_root() -> Path:
    from services.integration.tickets.rebar_config import load_rebar_backend_config

    return Path(load_rebar_backend_config().repo_root)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Import a Jira xlsx export into the local rebar ticket store."
    )
    parser.add_argument("xlsx", nargs="?", default="", help="Path to the Jira xlsx export.")
    parser.add_argument("--root", default="", help="Store path (defaults to the tickets config).")
    parser.add_argument(
        "--status",
        default="",
        help="Comma-separated Jira statuses to import (default: all rows).",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help="Import at most N tickets (0 means no limit). Use for a safe test batch.",
    )
    parser.add_argument(
        "--reset",
        action="store_true",
        help="Wipe the store and re-init empty with the reader identity before importing.",
    )
    args = parser.parse_args(argv if argv is not None else sys.argv[1:])
    repo_root = Path(args.root).expanduser() if args.root else _default_repo_root()

    if args.reset:
        reset_store(repo_root)
        print(f"[import-rebar-tickets] reset empty store at {repo_root.resolve()}")
        if not args.xlsx:
            return 0
    else:
        ensure_store(repo_root)

    if not args.xlsx:
        parser.error("an xlsx path is required unless only --reset is given")
    xlsx_path = Path(args.xlsx).expanduser()
    if not xlsx_path.is_file():
        parser.error(f"xlsx not found: {xlsx_path}")

    statuses = None
    if args.status:
        statuses = {s.strip().lower() for s in args.status.split(",") if s.strip()}

    report = import_xlsx(xlsx_path, repo_root, statuses=statuses, limit=args.limit)
    print(f"[import-rebar-tickets] store at {repo_root.resolve()}")
    print(
        f"[import-rebar-tickets] imported {len(report.created)}, "
        f"skipped {len(report.skipped)}, transitioned {report.transitioned}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
